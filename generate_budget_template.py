from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Code opération",
    "Libellé opération",
    "SDG",
    "Types d'achat",
    "Budget validé A (acquisitions)",
    "Budget initial",
    "Qtés validées A (à acquérir)",
    "Engagé (sans compter le RAR)",
    "Reste à engager",
    "Besoins + ou -",
    "Commentaire A",
    "BS",
    "Commentaire BS",
    "Nb Acquis",
    "DM",
    "Commentaire DM",
    "Montant acquis",
]

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

wb = Workbook()
ws = wb.active
ws.title = "Budget"

for index, name in enumerate(COLUMNS, start=1):
    cell = ws.cell(row=1, column=index, value=name)
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(index)].width = max(14, min(28, len(name) + 4))

ws.row_dimensions[1].height = 45
ws.freeze_panes = "C2"
ws.auto_filter.ref = "A1:" + get_column_letter(len(COLUMNS)) + "1"

COL_BUDGET = get_column_letter(COLUMNS.index("Budget validé A (acquisitions)") + 1)
COL_ENGAGE = get_column_letter(COLUMNS.index("Engagé (sans compter le RAR)") + 1)
COL_RESTE = get_column_letter(COLUMNS.index("Reste à engager") + 1)
COL_NB_ACQUIS = get_column_letter(COLUMNS.index("Nb Acquis") + 1)
COL_QTES = get_column_letter(COLUMNS.index("Qtés validées A (à acquérir)") + 1)
COL_BESOINS = get_column_letter(COLUMNS.index("Besoins + ou -") + 1)

for row in range(2, 502):
    reste_cell = ws[COL_RESTE + str(row)]
    reste_cell.value = (
        '=IF(' + COL_BUDGET + str(row) + '="","",' +
        COL_BUDGET + str(row) + '-' + COL_ENGAGE + str(row) + ')'
    )
    besoins_cell = ws[COL_BESOINS + str(row)]
    besoins_cell.value = (
        '=IF(' + COL_QTES + str(row) + '="","",' +
        COL_QTES + str(row) + '-' + COL_NB_ACQUIS + str(row) + ')'
    )

last_col = get_column_letter(len(COLUMNS))
for row in range(2, 502):
    for col in range(1, len(COLUMNS) + 1):
        ws.cell(row=row, column=col).border = THIN_BORDER
    ws[COL_BESOINS + str(row)].alignment = Alignment(horizontal="center")

wb.save("budget.xlsx")
print("OK")
